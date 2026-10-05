"""Write the staff report and the student workbook."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from shelf_report.logic import (
    ShelfItem,
    call_number_sort_key,
    excel_serial_to_text,
    normalize_barcode,
    parse_creator,
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, size=16, color="1F4E79")
BODY_FONT = Font(name="Calibri", size=11)
BOLD_FONT = Font(name="Calibri", bold=True, size=11)
NOTE_FILL = PatternFill("solid", fgColor="F8CBAD")
ISSUE_FILL = PatternFill("solid", fgColor="DDEBF7")
ORDER_FILL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="center")
HEADER_ALIGN = Alignment(wrap_text=True, vertical="center", horizontal="center")
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


def write_staff_report(
    path: Path,
    items: list[ShelfItem],
    missing: list[dict[str, str]],
    not_in_place: list[dict[str, str]],
    out_of_range: list[dict[str, str]],
    summary_rows: list[tuple[str, str]],
    detail_rows: list[tuple[str, str]],
) -> None:
    workbook = Workbook()
    shelf = workbook.active
    shelf.title = "Shelf read"
    _write_shelf_sheet(shelf, items)
    _write_alma_sheet(
        workbook.create_sheet("Missing items"),
        missing,
        [
            ("Shelf index", "_shelf_index"),
            ("Barcode", "Barcode"),
            ("Call number", "Call Number"),
            ("Description", "Description"),
            ("Copy ID", "Copy ID"),
            ("Title", "Title"),
            ("Creator", "_creator"),
            ("Permanent location", "Permanent Location"),
            ("Status", "Status"),
            ("Process type", "Process type"),
            ("On this shelf read", "_on_shelf_text"),
        ],
        sort_key=lambda record: call_number_sort_key(record.get("Call Number", "")),
    )
    _write_alma_sheet(
        workbook.create_sheet("On shelf but not in place"),
        not_in_place,
        [
            ("Shelf index", "_shelf_index"),
            ("Current location", "_current_location"),
            ("Barcode", "Barcode"),
            ("Call number", "Call Number"),
            ("Description", "Description"),
            ("Copy ID", "Copy ID"),
            ("Title", "Title"),
            ("Creator", "_creator"),
            ("Status", "Status"),
            ("Process type", "Process type"),
            ("At", "At"),
            ("Until", "_until"),
            ("Requests", "Requests"),
            ("Due date", "_due"),
            ("Permanent location", "Permanent Location"),
            ("Fulfillment note", "Fulfillment Note"),
            ("On this shelf read", "_on_shelf_text"),
        ],
    )
    _write_alma_sheet(
        workbook.create_sheet("Out of range"),
        out_of_range,
        [
            ("Shelf index", "_shelf_index"),
            ("Current location", "_current_location"),
            ("Barcode", "Barcode"),
            ("Call number", "Call Number"),
            ("Description", "Description"),
            ("Copy ID", "Copy ID"),
            ("Title", "Title"),
            ("Creator", "_creator"),
            ("Status", "Status"),
            ("Permanent location", "Permanent Location"),
            ("On this shelf read", "_on_shelf_text"),
        ],
    )
    _write_summary_sheet(workbook.create_sheet("Summary"), summary_rows, detail_rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def write_student_workbook(
    path: Path,
    items: list[ShelfItem],
    missing: list[dict[str, str]],
    search_locations: list[str],
    threshold: int,
    project_name: str,
) -> None:
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    cart_items = [item for item in items if item.barcode_not_found]
    order_items = [item for item in items if item.order_check]
    order_items.sort(key=lambda item: (item.shelf_index, item.barcode))
    _write_instructions(
        instructions,
        project_name,
        threshold,
        missing_count=len(missing),
        cart_count=len(cart_items),
        order_count=len(order_items),
    )
    _write_missing_student_sheet(
        workbook.create_sheet("01_Search for missing items"),
        missing,
        search_locations,
    )
    _write_cart_sheet(workbook.create_sheet("02_Put on a cart for Austin"), cart_items)
    _write_order_sheet(workbook.create_sheet("03_Possibly out of order items"), order_items)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def prepare_alma_records(
    records: list[dict[str, str]],
    items_by_barcode: dict[str, ShelfItem],
) -> list[dict[str, str]]:
    prepared = []
    for record in records:
        item = items_by_barcode.get(normalize_barcode(record.get("Barcode")))
        enriched = dict(record)
        enriched["_creator"] = parse_creator(record.get("Type / Creator / Imprint", ""))
        enriched["_until"] = excel_serial_to_text(record.get("Until", ""))
        enriched["_due"] = excel_serial_to_text(record.get("Due Date", ""))
        enriched["_on_shelf_text"] = "Yes" if item else "No"
        enriched["_shelf_index"] = "" if item is None else _excel_number(item.shelf_index)
        enriched["_current_location"] = "" if item is None else item.current_location
        prepared.append(enriched)
    return prepared


def _write_shelf_sheet(worksheet: Worksheet, items: list[ShelfItem]) -> None:
    headers = [
        "Index",
        "Barcode",
        "Notes",
        "Issue",
        "Current Location",
        "Order Check",
        "Analytic Index",
        "Normalized Call Number",
        "Permanent Call Number",
        "Description",
        "Item Copy Id",
        "Barcode",
        "Title",
        "Author",
        "Location Name",
    ]
    widths = [10, 18, 38, 28, 55, 55, 12, 26, 24, 28, 12, 18, 46, 32, 32]
    _write_header(worksheet, headers)
    for row_number, item in enumerate(items, start=2):
        values = [
            _excel_number(item.shelf_index),
            _excel_barcode(item.barcode),
            item.notes_text,
            item.issues_text,
            item.current_location,
            item.order_check,
            item.analytic_index if item.analytic_index is not None else "",
            item.normalized_call_number,
            item.permanent_call_number,
            item.description,
            item.copy_id,
            _excel_barcode(item.analytic_barcode) if item.analytic_barcode else "",
            item.title,
            item.author,
            item.location_name,
        ]
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row_number, column, value)
            cell.font = BODY_FONT
            cell.border = THIN
            cell.alignment = Alignment(vertical="center", wrap_text=column in {3, 4, 5, 6, 13, 14})
        if item.notes_text:
            worksheet.cell(row_number, 3).fill = NOTE_FILL
        if item.issues_text:
            worksheet.cell(row_number, 4).fill = ISSUE_FILL
        if item.order_check:
            worksheet.cell(row_number, 6).fill = ORDER_FILL
        worksheet.row_dimensions[row_number].height = 30
    _finish_table(worksheet, widths, len(items))
    worksheet.auto_filter.ref = f"A1:O{max(len(items) + 1, 2)}"
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:O{max(len(items), 1) + 1}"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_TABLOID
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.oddHeader.left.text = "Staff shelf read"
    worksheet.print_title_rows = "1:1"
    worksheet.page_setup.orientation = "landscape"
    worksheet.sheet_view.showGridLines = False


def _write_alma_sheet(worksheet, records, columns, sort_key=None) -> None:
    headers = [header for header, _key in columns]
    _write_header(worksheet, headers)
    ordered = sorted(records, key=sort_key) if sort_key else list(records)
    if not ordered:
        worksheet.cell(2, 1, "None in this run.").font = BODY_FONT
    for row_number, record in enumerate(ordered, start=2):
        for column, (_header, key) in enumerate(columns, start=1):
            value = record.get(key, "")
            if key == "Barcode":
                value = _excel_barcode(value)
            cell = worksheet.cell(row_number, column, value)
            cell.font = BODY_FONT
            cell.border = THIN
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        worksheet.row_dimensions[row_number].height = 30
    widths = [18 if "location" in header.lower() or header in {"Title", "Creator"} else 16 for header in headers]
    for header, width in zip(headers, widths):
        if header in {"Title", "Current location", "Fulfillment note"}:
            widths[headers.index(header)] = 42
        elif header in {"Call number", "Process type", "Permanent location"}:
            widths[headers.index(header)] = 28
    _finish_table(worksheet, widths, len(ordered))
    last_column = get_column_letter(len(headers))
    worksheet.auto_filter.ref = f"A1:{last_column}{max(len(ordered), 1) + 1}"
    worksheet.freeze_panes = "A2"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_title_rows = "1:1"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_TABLOID


def _write_summary_sheet(worksheet, summary_rows, detail_rows) -> None:
    worksheet["A1"] = "Shelf report summary"
    worksheet["A1"].font = TITLE_FONT
    worksheet.merge_cells("A1:B1")
    worksheet["A3"] = "Count"
    worksheet["B3"] = "Value"
    _style_header_cell(worksheet["A3"])
    _style_header_cell(worksheet["B3"])
    for offset, (label, value) in enumerate(summary_rows):
        worksheet.cell(4 + offset, 1, label).font = BODY_FONT
        worksheet.cell(4 + offset, 2, value).font = BODY_FONT
    detail_start = 6 + len(summary_rows)
    worksheet.cell(detail_start, 1, "Details").font = TITLE_FONT
    header_row = detail_start + 1
    worksheet.cell(header_row, 1, "Item")
    worksheet.cell(header_row, 2, "Detail")
    _style_header_cell(worksheet.cell(header_row, 1))
    _style_header_cell(worksheet.cell(header_row, 2))
    if not detail_rows:
        worksheet.cell(header_row + 1, 1, "Nothing else to review.").font = BODY_FONT
    for offset, (label, value) in enumerate(detail_rows):
        worksheet.cell(header_row + 1 + offset, 1, label).font = BODY_FONT
        cell = worksheet.cell(header_row + 1 + offset, 2, value)
        cell.font = BODY_FONT
        cell.alignment = WRAP
    worksheet.column_dimensions["A"].width = 46
    worksheet.column_dimensions["B"].width = 110
    worksheet.row_dimensions[1].height = 24
    worksheet.freeze_panes = "A4"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_LETTER


def _write_instructions(worksheet, project_name, threshold, missing_count, cart_count, order_count) -> None:
    lines = [
        (f"Shelf reading follow-up: {project_name}", TITLE_FONT),
        ("", BODY_FONT),
        ("Work through the numbered sheets in order. Mark the first column of each row as you finish it.", BODY_FONT),
        ("", BODY_FONT),
        (f"1. Search for missing items ({missing_count})", BOLD_FONT),
        (
            "These items were marked in place in Alma, but they were not on the shelf during this shelf read. "
            "Look in the Alma location first. Put an x in each other location column as you check it. "
            "If you find the item, write where it was in Found? (for example, Found in Rare Oversize).",
            BODY_FONT,
        ),
        ("", BODY_FONT),
        (f"2. Put on a cart for Austin ({cart_count})", BOLD_FONT),
        (
            "These barcodes were on the shelf, and Alma has no item record for them. "
            "Pull each item from the spot in Current Location and put it on the cart for Austin. "
            "Mark the first column once it is on the cart. If this sheet says there were none, skip it.",
            BODY_FONT,
        ),
        ("", BODY_FONT),
        (f"3. Possibly out of order items ({order_count})", BOLD_FONT),
        (
            f"These items are more than {threshold} positions away from where their call number would put them. "
            "The recommendation is approximate, especially for volumes, newspapers, and odd descriptions. "
            "Find the item at Current Location. Move it only when it is clearly out of order. "
            "Mark Rearranged? with Y, or write left if you leave it where it is.",
            BODY_FONT,
        ),
        ("", BODY_FONT),
        (
            "Items that are on the shelf but outside this call-number range are marked for staff. "
            "They are not on the out-of-order sheet, and they should stay where they are until staff retrieve them.",
            BODY_FONT,
        ),
    ]
    for row_number, (text, font) in enumerate(lines, start=1):
        cell = worksheet.cell(row_number, 1, text)
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        worksheet.row_dimensions[row_number].height = 36 if text and font is BODY_FONT else 20
    worksheet.column_dimensions["A"].width = 120
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 1
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_LETTER
    worksheet.sheet_view.showGridLines = False
    worksheet.print_options.horizontalCentered = True
    worksheet.page_margins.left = 0.6
    worksheet.page_margins.right = 0.6
    worksheet.page_margins.top = 0.6
    worksheet.page_margins.bottom = 0.6


def _write_missing_student_sheet(worksheet, records, search_locations) -> None:
    headers = [
        "Found?",
        "Call Number",
        "Description",
        "Copy Number",
        "Barcode",
        "Title",
        "Author",
        "Location from Alma (Check Here first)",
        "",
        *search_locations,
    ]
    if search_locations:
        worksheet.merge_cells(start_row=1, start_column=10, end_row=1, end_column=9 + len(search_locations))
        banner = worksheet.cell(1, 10, 'LOCATIONS TO CHECK (Put an "x" as you go)')
        banner.font = BOLD_FONT
        banner.alignment = Alignment(horizontal="center", vertical="center")
    _write_header(worksheet, headers, row_number=2)
    ordered = sorted(records, key=lambda record: call_number_sort_key(record.get("Call Number", "")))
    if not ordered:
        worksheet.cell(3, 1, "No missing items in this run.").font = BODY_FONT
    for row_number, record in enumerate(ordered, start=3):
        values = [
            "",
            record.get("Call Number", ""),
            record.get("Description", ""),
            record.get("Copy ID", ""),
            _excel_barcode(record.get("Barcode", "")),
            record.get("Title", ""),
            record.get("_creator", ""),
            record.get("Permanent Location", ""),
            "",
            *["" for _location in search_locations],
        ]
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row_number, column, value)
            cell.font = BODY_FONT
            cell.border = THIN
            cell.alignment = Alignment(vertical="center", wrap_text=column in {2, 6, 7, 8})
        worksheet.row_dimensions[row_number].height = 30
    widths = [24, 22, 24, 12, 18, 46, 32, 38, 3] + [18 for _location in search_locations]
    _finish_table(worksheet, widths, len(ordered), header_row=2)
    last_column = get_column_letter(len(headers))
    worksheet.freeze_panes = "A3"
    worksheet.auto_filter.ref = f"A2:{last_column}{max(len(ordered), 1) + 2}"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_TABLOID
    worksheet.print_title_rows = "1:2"
    worksheet.page_setup.horizontalCentered = True
    worksheet.oddHeader.left.text = "Search for missing items"


def _write_cart_sheet(worksheet, items: list[ShelfItem]) -> None:
    headers = [
        "Put on cart for Austin?",
        "Barcode from shelfread",
        "",
        "",
        "Current Location",
        "",
        "",
        "",
        "Permanent Call Number (if known)",
        "Description",
        "Copy Number",
        "",
        "Title",
        "Author",
    ]
    _write_header(worksheet, headers)
    if not items:
        worksheet.cell(2, 1, "No unmatched barcodes in this run.").font = BODY_FONT
    for row_number, item in enumerate(items, start=2):
        values = [
            "",
            _excel_barcode(item.barcode),
            "",
            "",
            item.current_location,
            "",
            "",
            "",
            item.permanent_call_number,
            item.description,
            item.copy_id,
            "",
            item.title,
            item.author,
        ]
        _write_body_row(worksheet, row_number, values, wrap_columns={5, 9, 13, 14})
    widths = [24, 22, 3, 3, 55, 3, 3, 3, 28, 24, 14, 3, 42, 28]
    _finish_table(worksheet, widths, len(items))
    _student_page_setup(worksheet, "Put on a cart for Austin")


def _write_order_sheet(worksheet, items: list[ShelfItem]) -> None:
    headers = [
        "Rearranged?",
        "Barcode",
        "",
        "",
        "Current Location",
        "What to do (recommendation is approximate):",
        "",
        "",
        "Permanent Call Number",
        "Description",
        "Copy Number",
        "",
        "Title",
        "Author",
    ]
    _write_header(worksheet, headers)
    if not items:
        worksheet.cell(2, 1, "No items were far enough out of order to list.").font = BODY_FONT
    for row_number, item in enumerate(items, start=2):
        values = [
            "",
            _excel_barcode(item.barcode),
            "",
            "",
            item.current_location,
            item.order_check,
            "",
            "",
            item.permanent_call_number,
            item.description,
            item.copy_id,
            "",
            item.title,
            item.author,
        ]
        _write_body_row(worksheet, row_number, values, wrap_columns={5, 6, 9, 13, 14})
        if item.order_check:
            worksheet.cell(row_number, 6).fill = ORDER_FILL
    widths = [16, 18, 3, 3, 48, 48, 3, 3, 24, 22, 14, 3, 40, 28]
    _finish_table(worksheet, widths, len(items))
    _student_page_setup(worksheet, "Possibly out of order")


def _student_page_setup(worksheet, header_text: str) -> None:
    last_column = get_column_letter(worksheet.max_column)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{last_column}{max(worksheet.max_row, 2)}"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_TABLOID
    worksheet.print_title_rows = "1:1"
    worksheet.oddHeader.left.text = header_text
    worksheet.page_setup.horizontalCentered = True


def _write_header(worksheet, headers, row_number: int = 1) -> None:
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(row_number, column, header)
        _style_header_cell(cell)
    worksheet.row_dimensions[row_number].height = 32


def _style_header_cell(cell) -> None:
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = HEADER_ALIGN


def _write_body_row(worksheet, row_number: int, values, wrap_columns: set[int]) -> None:
    for column, value in enumerate(values, start=1):
        cell = worksheet.cell(row_number, column, value)
        cell.font = BODY_FONT
        cell.border = THIN
        cell.alignment = Alignment(vertical="center", wrap_text=column in wrap_columns)
    worksheet.row_dimensions[row_number].height = 32


def _finish_table(worksheet, widths, row_count: int, header_row: int = 1) -> None:
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(widths))}{max(row_count, 1) + header_row}"
    worksheet.freeze_panes = f"A{header_row + 1}"
    worksheet.page_setup.orientation = "landscape"
    worksheet.sheet_view.zoomScale = 110
    worksheet.oddFooter.right.text = "Page &P of &N"
    worksheet.page_margins.left = 0.4
    worksheet.page_margins.right = 0.4
    worksheet.page_margins.top = 0.6
    worksheet.page_margins.bottom = 0.5
    worksheet.page_margins.header = 0.25
    worksheet.page_margins.footer = 0.25


def _excel_barcode(value: object):
    text = normalize_barcode(value)
    if text.isdigit() and len(text) <= 15:
        return int(text)
    return text


def _excel_number(value: object):
    number = float(value)
    if number.is_integer():
        return int(number)
    return number
